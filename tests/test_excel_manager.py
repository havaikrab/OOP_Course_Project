import os

import openpyxl

from src.excel_manager import ExcelManager
from src.vacancy import Vacancy


def test_excel_manager_first_init(excel_columns: dict) -> None:

    ExcelManager("test_data/some_test_data.xlsx")
    test_book = openpyxl.open("test_data/some_test_data.xlsx", read_only=True)
    test_sheet = test_book["vacancies"]
    for k, v in excel_columns.items():
        assert test_sheet[f"{k}1"].value == v
    test_book.close()
    os.remove("test_data/some_test_data.xlsx")


def test_excel_manager_no_sheet_init(excel_columns: dict) -> None:

    empty_book = openpyxl.Workbook()
    empty_book.save("test_data/some_test_data.xlsx")
    empty_book.close()

    vacancy_manager = ExcelManager("test_data/some_test_data.xlsx")
    test_book = openpyxl.open("test_data/some_test_data.xlsx", read_only=True)
    test_sheet = test_book["vacancies"]
    for k, v in excel_columns.items():
        assert test_sheet[f"{k}1"].value == v

    test_book.close()
    vacancies = vacancy_manager.read_data()
    assert vacancies == list()
    os.remove("test_data/some_test_data.xlsx")


def test_excel_manager_shared_resource(excel_columns: dict, saved_vacancies_data: list) -> None:

    empty_book = openpyxl.Workbook()
    empty_book.create_sheet("vacancies")
    empty_book.save("test_data/some_test_data.xlsx")
    empty_book.close()

    vacancy_manager = ExcelManager("test_data/some_test_data.xlsx")
    vacancies_list = [Vacancy(i) for i in saved_vacancies_data]
    vacancy_manager.update_data(vacancies_list)
    test_book = openpyxl.open("test_data/some_test_data.xlsx", read_only=True)
    test_sheet = test_book["vacancies"]
    for k, v in excel_columns.items():
        assert test_sheet[f"{k}1"].value == v

    for i in saved_vacancies_data:
        for k, v in excel_columns.items():
            if k < "J":
                assert test_sheet[f"{k}{saved_vacancies_data.index(i) + 2}"].value == i.get(v)
    test_book.close()

    other_manager = ExcelManager("test_data/some_test_data.xlsx")
    extracted_vacancies = other_manager.read_data()
    extracted_vacancies = other_manager.filter_by_salary_mode("за месяц", extracted_vacancies)
    extracted_vacancies = other_manager.sort_by_date(extracted_vacancies)
    created_list = [i.created_at for i in extracted_vacancies]
    assert created_list == [
        "2026-02-22T09:18:06+0300",
        "2026-02-20T15:06:57+0300",
        "2026-02-19T17:55:40+0300",
        "2026-02-18T14:57:32+0300",
        "2026-02-17T22:41:18+0300",
        "2026-02-17T20:02:18+0300",
        "2026-02-16T09:15:41+0300",
        "2026-02-04T09:48:06+0300",
        "2026-02-02T16:52:12+0300",
    ]

    other_manager.clear_data()
    assert other_manager.read_data() == list()

    os.remove("test_data/some_test_data.xlsx")
