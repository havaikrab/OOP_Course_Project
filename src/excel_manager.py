import openpyxl
from openpyxl import Workbook

from src.base_manager import BaseManager
from src.filters import MixinFilter
from src.utils import format_filename
from src.vacancy import Vacancy


class ExcelManager(BaseManager, MixinFilter):
    """Класс работы с Excel-файлами"""

    __columns = {
        "A": "hh_id",
        "B": "name",
        "C": "vacancy_link",
        "D": "location",
        "E": "created_at",
        "F": "employer_name",
        "G": "employer_link",
        "H": "requirements",
        "I": "responsibility",
        "J": "amount_from",
        "K": "amount_to",
        "L": "currency",
        "M": "mode",
        "N": "converted_from",
        "O": "converted_to",
        "P": "required_currency",
    }

    @classmethod
    def __describe_head(cls, empty_book: Workbook, filename: str) -> Workbook:
        """Приватный класс-метод, заполняющий заголовки столбцов в чистом листе Excel-файла"""

        empty_book.create_sheet("vacancies")
        data_sheet = empty_book["vacancies"]
        for k, v in ExcelManager.__columns.items():
            data_sheet[f"{k}1"] = v
        empty_book.save(filename)
        empty_book.close()
        return empty_book

    @classmethod
    def __check_existence(cls, filename: str) -> Workbook:
        """Приватный класс-метод, проверки существования Excel-файла и корректности имен его столбцов, содержащих
        информацию о вакансиях"""

        data_book: Workbook = openpyxl.Workbook()
        try:
            data_book = openpyxl.open(filename, read_only=False)
        except FileNotFoundError:
            data_book = openpyxl.Workbook()
            return cls.__describe_head(data_book, filename)
        else:
            try:
                data_sheet = data_book["vacancies"]
            except KeyError:
                return cls.__describe_head(data_book, filename)
            else:
                is_valid = True
                for k, v in ExcelManager.__columns.items():
                    if data_sheet[f"{k}1"].value != v:
                        is_valid = False
                        break
                if is_valid:
                    return data_book
                else:
                    return cls.__describe_head(data_book, filename)

    def __init__(self, filename: str = "data\\vacancies.xlsx"):
        """Метод инициализации объекта класса. Принимает строку с именем Excel-файла, который будет определен, как
        база данных класса"""

        self.__filename = format_filename(filename)
        self.__current_data: Workbook = ExcelManager.__check_existence(filename)
        self.__saved_id: set = set()

    def read_data(self) -> list[Vacancy]:
        """Метод, преобразующий содержимое файла базы данных в список объектов класса Vacancy и возвращающий его"""

        data_sheet = self.__current_data["vacancies"]
        result = list()
        for row in range(2, data_sheet.max_row + 1):
            result.append(Vacancy({v: data_sheet[f"{k}{row}"].value for k, v in ExcelManager.__columns.items()}))
        return result

    def clear_data(self) -> None:
        """Метод полной очистки файла базы данных"""

        data_book = openpyxl.Workbook()
        self.__current_data = ExcelManager.__describe_head(data_book, self.__filename)

    def __get_saved_id(self) -> None:
        """Приватный метод, обновляющий значение атрибута __saved_id, хранящий множество строк, соответствующих
        оригинальным номерам вакансий, размещенных на сайте hh.ru"""

        data_sheet = self.__current_data["vacancies"]
        self.__saved_id = set([data_sheet[f"A{row}"].value for row in range(2, data_sheet.max_row + 1)])

    def __exclude_duplicates(self, vacancies_list: list[Vacancy]) -> list:
        """Приватный метод, проверяющий передаваемый список на наличие вакансий, уже сохраненных в базе данных класса.
        Возвращает список, содержащий только новые вакансии"""

        self.__get_saved_id()
        result = list()
        result_id_set: set = set()
        for vacancy in vacancies_list:
            if vacancy.hh_id not in self.__saved_id and vacancy.hh_id not in result_id_set:
                result_id_set.add(vacancy.hh_id)
                result.append(vacancy)
        return result

    def update_data(self, some_data: list[Vacancy]) -> None:
        """Метод обновления базы данных класса передаваемым списком вакансий"""

        some_data = self.__exclude_duplicates(some_data)
        data_sheet = self.__current_data["vacancies"]
        row_number = data_sheet.max_row + 1
        for element in some_data:
            row_dict = {k: element.to_dict().get(v) for k, v in ExcelManager.__columns.items()}
            for k, v in row_dict.items():
                if v:
                    data_sheet[f"{k}{row_number}"] = v
            row_number += 1
        self.__current_data.save(self.__filename)
        self.__current_data.close()
